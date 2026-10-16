"""Construye el libro Excel. Usan formulas (COUNTIFS/SUMPRODUCT) sobre la hoja base."""
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .store import TZ

F = "Arial"
HDR = PatternFill("solid", fgColor="50C492") 


def _hdr(ws, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=F, bold=True, color="000000")
        cell.fill = HDR
        cell.alignment = Alignment(horizontal="center")


def build_workbook(rows, fecha, eventos) -> bytes:
    """rows: [(fecha_iso, hora_str, modulo, evento)] ordenadas."""
    wb = Workbook()
    base = wb.active
    base.title = "Base_Estructurada"
    _hdr(base, ["FECHA", "HORA", "NÚMERO DE MODULO"])
    for i, (f, h, m, _) in enumerate(rows, 2):
        base.cell(i, 1, date.fromisoformat(f)).number_format = "dd/mm/yyyy"
        base.cell(i, 2, time.fromisoformat(h)).number_format = "hh:mm:ss"
        base.cell(i, 3, m)
    last = max(len(rows) + 1, 2)
    base.freeze_panes = "A2"
    base.auto_filter.ref = f"A1:C{last}"
    for col, w in zip("ABC", (14, 12, 20)):
        base.column_dimensions[col].width = w
    mods = sorted({r[2] for r in rows})
    hours = sorted({int(r[1][:2]) for r in rows})
    rng_m = f"Base_Estructurada!$C$2:$C${last}"
    rng_h = f"Base_Estructurada!$B$2:$B${last}"
    bold = Font(name=F, bold=True)

    rm = wb.create_sheet("Resumen_por_Modulo")
    _hdr(rm, ["NÚMERO DE MODULO", "CANTIDAD DE REGISTROS"])
    for i, m in enumerate(mods, 2):
        rm.cell(i, 1, m)
        rm.cell(i, 2, f"=COUNTIFS({rng_m},A{i})")
    t = len(mods) + 2
    rm.cell(t, 1, "TOTAL").font = bold
    rm.cell(t, 2, f"=SUM(B2:B{t - 1})" if mods else 0).font = bold
    rm.column_dimensions["A"].width = 22
    rm.column_dimensions["B"].width = 26

    rh = wb.create_sheet("Resumen_por_Hora")
    _hdr(rh, ["HORA"] + [None] * len(mods) + ["TOTAL"])
    for j, m in enumerate(mods, 2):
        rh.cell(1, j, m).number_format = '"Módulo "0'
    for i, h in enumerate(hours, 2):
        rh.cell(i, 1, h).number_format = '00":00"'
        for j in range(2, len(mods) + 2):
            col = rh.cell(1, j).column_letter
            rh.cell(i, j, f"=SUMPRODUCT(({rng_m}={col}$1)*(HOUR({rng_h})=$A{i}))")
        endc = rh.cell(1, len(mods) + 1).column_letter
        rh.cell(i, len(mods) + 2, f"=SUM(B{i}:{endc}{i})")
    if hours:
        tr = len(hours) + 2
        rh.cell(tr, 1, "TOTAL").font = bold
        for j in range(2, len(mods) + 3):
            col = rh.cell(1, j).column_letter
            rh.cell(tr, j, f"=SUM({col}2:{col}{tr - 1})").font = bold
    rh.column_dimensions["A"].width = 12

    nt = wb.create_sheet("Notas")
    presentes = sorted({r[3] for r in rows})
    notas = [
        ("Generado", datetime.now(TZ).strftime("%d/%m/%Y %H:%M:%S")),
        ("Filtro de fecha", fecha or "Todas las fechas"),
        ("Eventos contados", ", ".join(eventos) if eventos else "Todos los tipos recibidos"),
        ("Tipos de evento en los datos", ", ".join(presentes) or "—"),
        ("Zona horaria", "America/Bogota (UTC-5); el origen entrega epoch UTC en milisegundos"),
        ("Deduplicación", "Por comando.id: cada mensaje llega duplicado con distinto event_id"),
        ("Aviso", "Si el tipo es healthcheck.check, el conteo corresponde a pulsos del módulo, "
                  "no a aperturas confirmadas de casillero. Ajuste COUNT_EVENTS cuando exista el evento de apertura."),
        ("Resúmenes", "Son fórmulas sobre Base_Estructurada: se recalculan si se edita la base."),
    ]
    for i, (k, v) in enumerate(notas, 1):
        nt.cell(i, 1, k).font = bold
        nt.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 30
    nt.column_dimensions["B"].width = 90

    for ws in wb.worksheets:  # Arial en todo el libro
        for row in ws.iter_rows():
            for c in row:
                if c.font.name != F:
                    c.font = Font(name=F, bold=c.font.bold, color=c.font.color)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
