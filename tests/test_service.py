import io
import json
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app

T0 = 1720306731662  # 06/07/2024 17:58:51 hora Colombia


def line(i, mod, ts, ev="healthcheck.check", event_id=None):
    body = {"comando": {"nombre": ev, "id": i, "timestamp": ts, "payload": {"moduloId": str(mod)}}, "validationStatus": 0}
    return (f"Sat Jul 06 17:58:51 GMT-05:00 2024 name=QUEUE_msg_received event_id=ID:{event_id or i} "
            f"msg_dest=Q msg_body={json.dumps(body)}")


def sample():
    # a y b llegan duplicadas con distinto event_id; c es del módulo 355 y otro tipo de evento
    return "\n".join([line("a", 342, T0, event_id="1"), line("a", 342, T0, event_id="2"),
                      line("b", 342, T0 + 30000), line("b", 342, T0 + 30000, event_id="x"),
                      line("c", 355, T0 + 60000, ev="locker.open"), "basura"])


def test_ingest_dedup_and_counts(tmp_path):
    c = TestClient(create_app(db_path=str(tmp_path / "t.db")))
    r = c.post("/ingest", content=sample()).json()
    assert r == {"recibidas": 6, "nuevas": 3, "duplicadas": 2, "invalidas": 1}
    assert c.post("/ingest", content=sample()).json()["nuevas"] == 0  # idempotente
    assert c.get("/counts?fecha=2024-07-06").json()["modulos"] == {"342": 2, "355": 1}
    assert c.get("/counts?fecha=2024-07-06&evento=locker.open").json()["modulos"] == {"355": 1}
    assert c.get("/counts?fecha=2024-07-07").json()["total"] == 0
    assert c.get("/counts?fecha=mal").status_code == 422


def test_export_excel(tmp_path):
    c = TestClient(create_app(db_path=str(tmp_path / "t.db")))
    c.post("/ingest", content=sample())
    r = c.get("/export.xlsx?fecha=2024-07-06&modulo=342")
    ws = load_workbook(io.BytesIO(r.content))["Base_Estructurada"]
    assert [x.value for x in ws[1]] == ["FECHA", "HORA", "NÚMERO DE MODULO"]
    assert ws.max_row == 3 and ws["C2"].value == 342 and str(ws["B2"].value) == "17:58:51"


def test_tail_real_time(tmp_path):
    log = tmp_path / "log.txt"
    log.write_text(line("a", 342, T0) + "\n")
    with TestClient(create_app(db_path=str(tmp_path / "t.db"), log_path=str(log))) as c:
        def total():
            return c.get("/counts?fecha=all").json()["total"]
        for _ in range(40):
            if total() == 1: break
            time.sleep(0.1)
        assert total() == 1
        with open(log, "a") as f:
            f.write(line("b", 355, T0 + 30000) + "\n")
        for _ in range(40):
            if total() == 2: break
            time.sleep(0.1)
        assert total() == 2
